from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from lxml import etree

# -----------------------------------------------------------------------
# COLOR PALETTE
# Orange & Gold = PRIMARY  |  White = Backgrounds  |  Red = Highlights only
# -----------------------------------------------------------------------
ORANGE      = RGBColor(0xE8, 0x60, 0x1A)   # ★ Primary – warm orange (most used)
GOLD        = RGBColor(0xF4, 0xA8, 0x00)   # ★ Secondary – golden yellow (2nd most used)
DARK_ORANGE = RGBColor(0xB8, 0x40, 0x00)   # Dark shade of orange (header bars, dark bg)
RED_ACCENT  = RGBColor(0xC0, 0x00, 0x00)   # ✦ Accent only – small highlights, dividers
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
CREAM       = RGBColor(0xFF, 0xF8, 0xF0)   # Warm white slide background
LIGHT_BG    = RGBColor(0xFD, 0xF0, 0xE0)   # Very light orange tint for sections
DARK_TEXT   = RGBColor(0x2A, 0x18, 0x00)   # Near-black warm text (readable on white/cream)
MID_TEXT    = RGBColor(0x5A, 0x38, 0x00)   # Mid-brown text for secondary info

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

# -----------------------------------------------------------------------
# HELPERS
# -----------------------------------------------------------------------
def set_bg(slide, rgb):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb

def rect(slide, l, t, w, h, fill=None, line=None, lw=0):
    s = slide.shapes.add_shape(1, l, t, w, h)
    s.line.width = Pt(lw) if line else Pt(0)
    if line:
        s.line.color.rgb = line
    else:
        s.line.fill.background()
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    return s

def tb(slide, l, t, w, h, text, size, bold=False, color=None,
       align=PP_ALIGN.LEFT, italic=False):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color:
        r.font.color.rgb = color
    return box

def divider(slide, l, t, w, color=RED_ACCENT, h_pt=2):
    s = slide.shapes.add_shape(1, l, t, w, Pt(h_pt))
    s.line.fill.background()
    s.fill.solid()
    s.fill.fore_color.rgb = color
    return s

def clear_slide(slide):
    tree = slide.shapes._spTree
    for sp in list(tree):
        tag = sp.tag.split('}')[-1]
        if tag in ('sp', 'pic', 'cxnSp', 'graphicFrame', 'grpSp'):
            tree.remove(sp)

def header(slide, title, sub=None):
    """Orange gradient-style banner header for content slides."""
    # Main banner – dark orange
    rect(slide, 0, 0, SLIDE_W, Inches(1.38), fill=DARK_ORANGE)
    # Gold accent left strip
    rect(slide, 0, 0, Inches(0.18), Inches(1.38), fill=GOLD)
    # Small red right-edge accent
    rect(slide, SLIDE_W - Inches(0.08), 0, Inches(0.08), Inches(1.38), fill=RED_ACCENT)
    # Title text
    tb(slide, Inches(0.35), Inches(0.18), Inches(12.5), Inches(0.75),
       title, 30, bold=True, color=WHITE)
    if sub:
        tb(slide, Inches(0.35), Inches(0.92), Inches(12.5), Inches(0.38),
           sub, 12, color=GOLD, align=PP_ALIGN.LEFT)

def footer(slide):
    # Gold footer bar
    rect(slide, 0, SLIDE_H - Inches(0.14), SLIDE_W, Inches(0.14), fill=GOLD)
    # Tiny red edge
    rect(slide, 0, SLIDE_H - Inches(0.14), Inches(0.35), Inches(0.14), fill=RED_ACCENT)

# -----------------------------------------------------------------------
# BUILD
# -----------------------------------------------------------------------
def build(template, output):
    prs = Presentation(template)
    slides = prs.slides

    # ==============================================================
    # SLIDE 1 – Title
    # ==============================================================
    sl = slides[0]; clear_slide(sl)
    set_bg(sl, WHITE)

    # Top banner (orange)
    rect(sl, 0, 0, SLIDE_W, Inches(1.15), fill=ORANGE)
    rect(sl, 0, 0, Inches(0.22), Inches(1.15), fill=GOLD)
    rect(sl, SLIDE_W - Inches(0.22), 0, Inches(0.22), Inches(1.15), fill=DARK_ORANGE)
    tb(sl, Inches(0.35), Inches(0.1), Inches(12.6), Inches(0.55),
       "Indore Institute of Science & Technology",
       22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    tb(sl, Inches(0.35), Inches(0.65), Inches(12.6), Inches(0.38),
       "Department of Information Technology",
       14, color=LIGHT_BG, align=PP_ALIGN.CENTER)

    # "A Minor Project…" label
    tb(sl, Inches(1), Inches(1.28), Inches(11.3), Inches(0.42),
       "A Minor Project Presentation on",
       14, color=MID_TEXT, align=PP_ALIGN.CENTER)

    # Title box – dark orange bg, gold+white text
    rect(sl, Inches(1.4), Inches(1.88), Inches(10.5), Inches(1.72), fill=DARK_ORANGE)
    rect(sl, Inches(1.4), Inches(1.88), Inches(0.1), Inches(1.72), fill=GOLD)
    rect(sl, Inches(11.8), Inches(1.88), Inches(0.1), Inches(1.72), fill=GOLD)
    tb(sl, Inches(1.6), Inches(1.98), Inches(10.1), Inches(0.65),
       "CYBERTRACE", 36, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    tb(sl, Inches(1.6), Inches(2.68), Inches(10.1), Inches(0.78),
       "Real-Time Network Security Monitoring  &  Anomaly Detection System",
       17, color=WHITE, align=PP_ALIGN.CENTER)

    # Divider line
    divider(sl, Inches(1.3), Inches(3.78), Inches(10.7), color=ORANGE, h_pt=2)

    # Guided by / Presented by
    tb(sl, Inches(1.0), Inches(3.98), Inches(4.5), Inches(0.42),
       "Guided By:", 14, bold=True, color=DARK_ORANGE)
    tb(sl, Inches(1.0), Inches(4.42), Inches(4.5), Inches(0.78),
       "[Professor Name]\nDepartment of IT", 13, color=DARK_TEXT)

    tb(sl, Inches(7.5), Inches(3.98), Inches(4.5), Inches(0.42),
       "Presented By:", 14, bold=True, color=DARK_ORANGE)
    tb(sl, Inches(7.5), Inches(4.42), Inches(4.5), Inches(0.78),
       "Priyanshi Gupta  [Enrollment No.]\n[Team Member Names]", 13, color=DARK_TEXT)

    # Year
    tb(sl, Inches(3.5), Inches(5.45), Inches(6.0), Inches(0.42),
       "Academic Year: 2024 – 2025", 13, color=MID_TEXT, align=PP_ALIGN.CENTER)

    footer(sl)

    # ==============================================================
    # SLIDE 2 – Outline
    # ==============================================================
    sl = slides[1]; clear_slide(sl)
    set_bg(sl, CREAM)
    header(sl, "Outline", "Topics Covered in This Presentation")
    footer(sl)

    topics = [
        ("01", "Proposed Area & Topic"),
        ("02", "Introduction"),
        ("03", "Objectives"),
        ("04", "Proposed Methodology"),
        ("05", "Technical Details"),
        ("06", "Expected Outcome"),
        ("07", "References"),
    ]
    box_w, box_h = Inches(5.5), Inches(0.64)
    cols_x = [Inches(0.7), Inches(7.0)]
    base_y = Inches(1.62)

    for idx, (num, name) in enumerate(topics):
        cx = cols_x[idx % 2]
        cy = base_y + (idx // 2) * Inches(0.75)
        # Alternating: orange / gold background
        bg = ORANGE if idx % 2 == 0 else GOLD
        txt = WHITE if idx % 2 == 0 else DARK_TEXT
        rect(sl, cx, cy, box_w, box_h, fill=bg)
        # Number badge – white box
        rect(sl, cx, cy, Inches(0.58), box_h, fill=WHITE)
        tb(sl, cx + Inches(0.04), cy + Inches(0.1), Inches(0.5), Inches(0.45),
           num, 18, bold=True, color=ORANGE if idx % 2 == 0 else GOLD,
           align=PP_ALIGN.CENTER)
        tb(sl, cx + Inches(0.66), cy + Inches(0.12), Inches(4.7), Inches(0.42),
           name, 17, bold=True, color=txt)

    # ==============================================================
    # SLIDE 3 – Area & Topic
    # ==============================================================
    sl = slides[2]; clear_slide(sl)
    set_bg(sl, CREAM)
    header(sl, "Proposed Area & Topic", "Problem Domain  |  Research Focus  |  Motivation")
    footer(sl)

    # Left card – light warm bg
    rect(sl, Inches(0.5), Inches(1.5), Inches(5.85), Inches(5.65), fill=LIGHT_BG)
    rect(sl, Inches(0.5), Inches(1.5), Inches(5.85), Inches(0.52), fill=ORANGE)
    tb(sl, Inches(0.62), Inches(1.55), Inches(5.5), Inches(0.4),
       "AREA", 14, bold=True, color=WHITE)

    for i, pt in enumerate([
        "• Cybersecurity & Network Traffic Analysis",
        "• Intrusion Detection Systems (IDS)",
        "• Real-time Packet Monitoring",
    ]):
        tb(sl, Inches(0.65), Inches(2.18) + i * Inches(0.65), Inches(5.5), Inches(0.55),
           pt, 14, color=DARK_TEXT)

    # Mini divider
    divider(sl, Inches(0.65), Inches(3.85), Inches(5.4), color=ORANGE, h_pt=1.5)

    tb(sl, Inches(0.65), Inches(4.0), Inches(5.5), Inches(0.38),
       "MOTIVATION", 13, bold=True, color=RED_ACCENT)
    tb(sl, Inches(0.65), Inches(4.42), Inches(5.5), Inches(1.55),
       "Cyber attacks are rising exponentially. Traditional monitoring tools are complex and siloed. CyberTrace brings clarity to threat detection using open-source technology.",
       13, color=DARK_TEXT)

    # Right card – dark orange bg
    rect(sl, Inches(6.6), Inches(1.5), Inches(6.25), Inches(5.65), fill=DARK_ORANGE)
    rect(sl, Inches(6.6), Inches(1.5), Inches(6.25), Inches(0.52), fill=GOLD)
    tb(sl, Inches(6.72), Inches(1.55), Inches(6.0), Inches(0.4),
       "TOPIC", 14, bold=True, color=DARK_TEXT)
    tb(sl, Inches(6.72), Inches(2.15), Inches(5.85), Inches(0.55),
       "CyberTrace", 26, bold=True, color=GOLD)
    tb(sl, Inches(6.72), Inches(2.78), Inches(5.85), Inches(0.72),
       "Real-Time Network Security Monitoring\n& Anomaly Detection System",
       15, color=WHITE)
    divider(sl, Inches(6.72), Inches(3.65), Inches(5.6), color=GOLD, h_pt=1.5)
    tb(sl, Inches(6.72), Inches(3.82), Inches(5.85), Inches(3.1),
       "The project captures live or recorded network packets, passes them through anomaly detection engines (rule-based + ML thresholding), and displays real-time metrics on a modern web dashboard built with React + FastAPI.",
       14, color=LIGHT_BG)

    # ==============================================================
    # SLIDE 4 – Introduction
    # ==============================================================
    sl = slides[3]; clear_slide(sl)
    set_bg(sl, CREAM)
    header(sl, "Introduction", "What is CyberTrace?  |  Why does it matter?")
    footer(sl)

    # Quote banner – gold bg
    rect(sl, Inches(0.5), Inches(1.52), Inches(12.3), Inches(0.88), fill=GOLD)
    rect(sl, Inches(0.5), Inches(1.52), Inches(0.12), Inches(0.88), fill=RED_ACCENT)
    tb(sl, Inches(0.75), Inches(1.66), Inches(11.85), Inches(0.6),
       '"CyberTrace bridges raw network data and actionable security intelligence — in real time."',
       16, bold=True, color=DARK_TEXT, align=PP_ALIGN.CENTER, italic=True)

    cards = [
        ("What It Does",   "Captures & parses live or PCAP-based network packets and runs them through multiple detection algorithms to flag threats automatically."),
        ("How It Works",   "A decoupled architecture: Scapy handles capture, FastAPI serves a REST API, and React renders beautiful, real-time dashboards."),
        ("Why It Matters", "Early threat detection reduces incident response time. CyberTrace makes that accessible without enterprise cost."),
    ]
    hdr_colors = [ORANGE, DARK_ORANGE, ORANGE]
    card_w = Inches(3.9)
    for (title, body), cx, hc in zip(cards, [Inches(0.5), Inches(4.72), Inches(8.94)], hdr_colors):
        rect(sl, cx, Inches(2.58), card_w, Inches(4.45), fill=WHITE,
             line=GOLD, lw=0.8)
        rect(sl, cx, Inches(2.58), card_w, Inches(0.56), fill=hc)
        tb(sl, cx + Inches(0.12), Inches(2.64), card_w - Inches(0.22), Inches(0.44),
           title, 14, bold=True, color=WHITE)
        tb(sl, cx + Inches(0.15), Inches(3.27), card_w - Inches(0.3), Inches(3.65),
           body, 13, color=DARK_TEXT)

    # ==============================================================
    # SLIDE 5 – Objectives
    # ==============================================================
    sl = slides[4]; clear_slide(sl)
    set_bg(sl, CREAM)
    header(sl, "Objectives", "Clear Goals for the CyberTrace Project")
    footer(sl)

    objectives = [
        ("01", "Packet Capture",     "Efficiently capture and parse live or file-based network traffic using Scapy."),
        ("02", "Anomaly Detection",  "Detect threats like DoS attacks, Port Scans, and DNS anomalies in real time."),
        ("03", "Web Dashboard",      "Provide an intuitive, responsive UI to visualize traffic metrics and alerts."),
        ("04", "PCAP Forensics",     "Allow post-incident analysis by uploading .pcap files for offline investigation."),
        ("05", "Decoupled Design",   "Implement a modern, scalable architecture (FastAPI backend + React frontend)."),
    ]
    obj_w, obj_h = Inches(12.0), Inches(0.77)
    # Alternating: orange row / light-bg row
    for i, (num, title, desc) in enumerate(objectives):
        y = Inches(1.52) + i * (obj_h + Inches(0.08))
        if i % 2 == 0:
            bg, num_bg, title_col, desc_col = ORANGE, GOLD, WHITE, WHITE
        else:
            bg, num_bg, title_col, desc_col = LIGHT_BG, ORANGE, DARK_TEXT, DARK_TEXT
        rect(sl, Inches(0.65), y, obj_w, obj_h, fill=bg)
        rect(sl, Inches(0.65), y, Inches(0.62), obj_h, fill=num_bg)
        tb(sl, Inches(0.67), y + Inches(0.14), Inches(0.58), Inches(0.5),
           num, 16, bold=True, color=WHITE if i % 2 == 0 else DARK_TEXT,
           align=PP_ALIGN.CENTER)
        tb(sl, Inches(1.4), y + Inches(0.1), Inches(2.9), Inches(0.5),
           title, 15, bold=True, color=title_col)
        tb(sl, Inches(4.45), y + Inches(0.1), Inches(8.0), Inches(0.5),
           desc, 13, color=desc_col)

    # ==============================================================
    # SLIDE 6 – Methodology
    # ==============================================================
    sl = slides[5]; clear_slide(sl)
    set_bg(sl, CREAM)
    header(sl, "Proposed Methodology", "System Architecture  |  Data Flow  |  Algorithm")
    footer(sl)

    layers = [
        ("Layer 1", "Data Capture",    "Scapy sniffs packets from network interface OR reads .pcap files via PCAPReader class.",   ORANGE),
        ("Layer 2", "Detection Engine","Each packet passes through DoSDetector, PortScanDetector, DNSAnomalyDetector, MLAnomalyDetector.", GOLD),
        ("Layer 3", "Backend API",     "FastAPI exposes /api/metrics, /api/alerts, /api/upload_pcap endpoints. WebSocket for live data.", DARK_ORANGE),
        ("Layer 4", "Frontend UI",     "React + Vite renders dashboards with charts, alert panels, and PCAP Forensics upload page.", ORANGE),
    ]

    for i, (num, title, desc, col) in enumerate(layers):
        y = Inches(1.52) + i * Inches(1.28)
        rect(sl, Inches(0.5), y, Inches(5.8), Inches(1.08), fill=WHITE,
             line=col, lw=0.6)
        rect(sl, Inches(0.5), y, Inches(1.12), Inches(1.08), fill=col)
        tb(sl, Inches(0.52), y + Inches(0.06), Inches(1.08), Inches(0.44),
           num, 11, bold=True, color=WHITE if col != GOLD else DARK_TEXT,
           align=PP_ALIGN.CENTER)
        tb(sl, Inches(0.52), y + Inches(0.54), Inches(1.08), Inches(0.44),
           title, 10, bold=True, color=WHITE if col != GOLD else DARK_TEXT,
           align=PP_ALIGN.CENTER)
        tb(sl, Inches(1.74), y + Inches(0.12), Inches(4.4), Inches(0.88),
           desc, 12, color=DARK_TEXT)
        if i < 3:
            tb(sl, Inches(5.9), y + Inches(0.58), Inches(0.8), Inches(0.4),
               "▼", 18, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)

    # Right side – flow summary, dark-orange bg
    rect(sl, Inches(6.72), Inches(1.52), Inches(6.0), Inches(5.15), fill=DARK_ORANGE)
    rect(sl, Inches(6.72), Inches(1.52), Inches(6.0), Inches(0.52), fill=GOLD)
    tb(sl, Inches(6.84), Inches(1.57), Inches(5.75), Inches(0.4),
       "Data Flow Summary", 14, bold=True, color=DARK_TEXT)
    flow = (
        "Network Interface  ➜  Scapy Sniffer\n"
        "         ⬇\n"
        "Packet Parser (parse_packet)\n"
        "         ⬇\n"
        "Anomaly Detectors (4 Engines)\n"
        "         ⬇\n"
        "AlertManager  +  Database\n"
        "         ⬇\n"
        "FastAPI REST Endpoints\n"
        "         ⬇\n"
        "React Dashboard (Port 5173)"
    )
    tb(sl, Inches(6.84), Inches(2.18), Inches(5.75), Inches(4.2), flow, 13, color=WHITE)

    # ==============================================================
    # SLIDE 7 – Technical Details
    # ==============================================================
    sl = slides[6]; clear_slide(sl)
    set_bg(sl, CREAM)
    header(sl, "Technical Details", "Hardware & Software Requirements  |  Tech Stack")
    footer(sl)

    # Hardware – light bg, orange header
    rect(sl, Inches(0.5), Inches(1.52), Inches(5.72), Inches(5.55), fill=LIGHT_BG)
    rect(sl, Inches(0.5), Inches(1.52), Inches(5.72), Inches(0.56), fill=ORANGE)
    tb(sl, Inches(0.62), Inches(1.57), Inches(5.5), Inches(0.44),
       "⚙  Hardware Requirements", 14, bold=True, color=WHITE)
    for i, item in enumerate([
        "• Standard PC / Laptop",
        "• Minimum 4 GB RAM (8 GB Recommended)",
        "• Multi-core Processor (Intel i3+ / AMD Ryzen 3+)",
        "• Network Interface Card (NIC) for live capture",
        "• 500 MB free disk space",
    ]):
        tb(sl, Inches(0.68), Inches(2.25) + i * Inches(0.67), Inches(5.3), Inches(0.58),
           item, 13, color=DARK_TEXT)

    # Software – dark orange bg, gold header
    rect(sl, Inches(6.62), Inches(1.52), Inches(6.25), Inches(5.55), fill=DARK_ORANGE)
    rect(sl, Inches(6.62), Inches(1.52), Inches(6.25), Inches(0.56), fill=GOLD)
    tb(sl, Inches(6.74), Inches(1.57), Inches(6.0), Inches(0.44),
       "💻  Tech Stack", 14, bold=True, color=DARK_TEXT)
    sw = [
        ("Language",   "Python 3.x  |  JavaScript (ES2022)"),
        ("Capture",    "Scapy (Packet Sniffing & PCAP Parsing)"),
        ("Backend",    "FastAPI + Uvicorn (REST API, WebSocket)"),
        ("Frontend",   "React JS + Vite + TailwindCSS + DaisyUI"),
        ("Charts",     "Recharts (Interactive Data Visualizations)"),
        ("Database",   "SQLite (via SQLAlchemy ORM)"),
        ("OS Support", "Linux (Recommended)  |  Windows"),
    ]
    for i, (key, val) in enumerate(sw):
        y = Inches(2.24) + i * Inches(0.64)
        tb(sl, Inches(6.74), y, Inches(2.1), Inches(0.5),
           key + ":", 12, bold=True, color=GOLD)
        tb(sl, Inches(8.98), y, Inches(3.72), Inches(0.5),
           val, 12, color=WHITE)

    # ==============================================================
    # SLIDE 8 – Expected Outcome
    # ==============================================================
    sl = slides[7]; clear_slide(sl)
    set_bg(sl, CREAM)
    header(sl, "Expected Outcome", "Deliverables  |  Results  |  Impact")
    footer(sl)

    outcomes = [
        ("🛡", "Functional Prototype",    "A working system capable of real-time packet sniffing, anomaly detection, and alert generation."),
        ("📊", "Modern Web Dashboard",    "A premium React-based UI with live metrics, chart visualizations, and a threat alert panel."),
        ("🔍", "PCAP Forensics Module",   "Post-incident analysis by uploading .pcap files to extract packets and generate threat reports."),
        ("🏗", "Scalable Architecture",   "A fully decoupled FastAPI + React pattern that can be easily extended for production use."),
    ]
    hdr_c = [ORANGE, GOLD, DARK_ORANGE, ORANGE]
    title_c = [WHITE, DARK_TEXT, WHITE, WHITE]
    positions = [
        (Inches(0.5),  Inches(1.52)),
        (Inches(6.92), Inches(1.52)),
        (Inches(0.5),  Inches(4.28)),
        (Inches(6.92), Inches(4.28)),
    ]
    cw = Inches(5.78)
    for (icon, title, desc), (cx, cy), hc, tc in zip(outcomes, positions, hdr_c, title_c):
        rect(sl, cx, cy, cw, Inches(2.32), fill=WHITE, line=GOLD, lw=0.6)
        rect(sl, cx, cy, cw, Inches(0.58), fill=hc)
        tb(sl, cx + Inches(0.12), cy + Inches(0.06), Inches(0.55), Inches(0.46),
           icon, 22, align=PP_ALIGN.CENTER)
        tb(sl, cx + Inches(0.78), cy + Inches(0.1), Inches(4.75), Inches(0.42),
           title, 14, bold=True, color=tc)
        tb(sl, cx + Inches(0.15), cy + Inches(0.72), Inches(5.5), Inches(1.48),
           desc, 13, color=DARK_TEXT)

    # ==============================================================
    # SLIDE 9 – References
    # ==============================================================
    sl = slides[8]; clear_slide(sl)
    set_bg(sl, CREAM)
    header(sl, "References", "")
    footer(sl)

    refs = [
        "[1]  RFC 793 – Transmission Control Protocol (TCP). IETF, 1981.",
        "[2]  Scapy Documentation – Packet Manipulation Library. https://scapy.net",
        "[3]  FastAPI Official Documentation – High Performance APIs with Python. https://fastapi.tiangolo.com",
        "[4]  React JS Official Documentation – A JavaScript Library for Building UIs. https://react.dev",
        '[5]  V. Chandola et al. – "Anomaly Detection: A Survey". ACM Computing Surveys, 2009.',
        "[6]  TailwindCSS & DaisyUI Component Library – https://daisyui.com",
    ]
    rect(sl, Inches(0.5), Inches(1.52), Inches(12.35), Inches(5.6), fill=LIGHT_BG)
    side_colors = [ORANGE, GOLD, ORANGE, GOLD, ORANGE, GOLD]
    for i, ref in enumerate(refs):
        y = Inches(1.72) + i * Inches(0.84)
        rect(sl, Inches(0.62), y, Inches(12.1), Inches(0.74), fill=WHITE)
        rect(sl, Inches(0.62), y, Inches(0.08), Inches(0.74), fill=side_colors[i])
        tb(sl, Inches(0.84), y + Inches(0.12), Inches(11.75), Inches(0.52),
           ref, 13, color=DARK_TEXT)

    prs.save(output)
    print(f"✅  Saved → {output}")


if __name__ == "__main__":
    build(
        "/home/pinkparrot/Priyanshi Gupta/format.pptx",
        "/home/pinkparrot/Priyanshi Gupta/CyberTrace/CyberTrace_Viva_Presentation.pptx"
    )
