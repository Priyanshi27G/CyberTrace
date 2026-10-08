import sys
from pptx import Presentation

def extract_text(filename):
    try:
        prs = Presentation(filename)
        for i, slide in enumerate(prs.slides):
            print(f"\n--- Slide {i+1} ---")
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    print(shape.text)
    except Exception as e:
        print(f"Error reading {filename}: {e}")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        extract_text(sys.argv[1])
    else:
        print("Usage: python read_pptx.py <filename>")
